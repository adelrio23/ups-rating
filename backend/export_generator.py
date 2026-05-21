"""
Export generator: produces customer-facing and internal Excel workbooks,
plus a Word deal brief.
"""
import io
from typing import Optional
import xlsxwriter
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ZONES = list(range(2, 9))
GREEN = "#C6EFCE"
GREEN_FONT = "#276221"
GREY = "#D9D9D9"
GREY_FONT = "#595959"
BLUE = "#BDD7EE"
BLUE_FONT = "#1F497D"
RED = "#FFC7CE"
RED_FONT = "#9C0006"
HEADER_BG = "#2F5496"
HEADER_FONT = "#FFFFFF"


def _hex_to_rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


def generate_customer_excel(scenario_result: dict, scenario_name: str = "Proposal") -> bytes:
    output = io.BytesIO()
    wb = xlsxwriter.Workbook(output, {"in_memory": True})

    summary = scenario_result["summary"]
    packages = scenario_result["packages"]
    rsm = scenario_result.get("rate_source_maps", {})

    # Formats
    header_fmt = wb.add_format({
        "bold": True, "bg_color": HEADER_BG, "font_color": HEADER_FONT,
        "border": 1, "align": "center", "valign": "vcenter",
    })
    money_fmt = wb.add_format({"num_format": "$#,##0.00", "border": 1, "align": "right"})
    pct_fmt = wb.add_format({"num_format": "0.0%", "border": 1, "align": "right"})
    green_fmt = wb.add_format({
        "bg_color": GREEN, "font_color": GREEN_FONT,
        "num_format": "$#,##0.00", "border": 1, "align": "right",
    })
    grey_fmt = wb.add_format({
        "bg_color": GREY, "font_color": GREY_FONT,
        "num_format": "$#,##0.00", "border": 1, "align": "right",
    })
    blue_fmt = wb.add_format({
        "bg_color": BLUE, "font_color": BLUE_FONT,
        "num_format": "$#,##0.00", "border": 1, "align": "right",
    })
    dash_fmt = wb.add_format({
        "bg_color": BLUE, "font_color": BLUE_FONT,
        "border": 1, "align": "center",
    })
    title_fmt = wb.add_format({"bold": True, "font_size": 14})
    label_fmt = wb.add_format({"bold": True})
    plain_fmt = wb.add_format({"border": 1, "align": "center"})

    # ── OVERVIEW TAB ────────────────────────────────────────────────────────────
    ws = wb.add_worksheet("Overview")
    ws.set_column("A:A", 30)
    ws.set_column("B:B", 20)
    ws.write("A1", f"Pricing Proposal: {scenario_name}", title_fmt)

    rows = [
        ("Monthly Packages", summary.get("total_packages", 0)),
        ("Customer Current Monthly Spend", f'${summary.get("customer_total_monthly", 0):,.2f}'),
        ("Proposed Monthly Spend", f'${summary.get("sell_total_monthly", 0):,.2f}'),
        ("Customer Monthly Savings", f'${summary.get("customer_savings_monthly", 0):,.2f}'),
        ("Odyssey Monthly Margin", f'${summary.get("margin_monthly", 0):,.2f}'),
        ("Margin %", f'{summary.get("margin_pct", 0):.1f}%'),
        ("Markup Applied", f'{summary.get("markup_pct", 0):.1f}%'),
        ("Our Fuel %", f'{summary.get("our_fuel_pct", 0):.1f}%'),
        ("Customer Fuel %", f'{summary.get("customer_fuel_pct", 0):.1f}%'),
        ("Fuel Spread", f'{summary.get("fuel_spread_pct", 0):.1f}%'),
        ("Green Packages (saves)", summary.get("green_count", 0)),
        ("Grey Packages (match)", summary.get("grey_count", 0)),
        ("Red Packages (no win)", summary.get("red_count", 0)),
        ("Corrections (excluded)", summary.get("corrections_count", 0)),
    ]
    for i, (label, val) in enumerate(rows, start=2):
        ws.write(f"A{i}", label, label_fmt)
        ws.write(f"B{i}", val)

    # ── RATE GRID TABS ────────────────────────────────────────────────────────
    for svc_key, svc_label in [
        ("ground_commercial", "Ground Commercial"),
        ("ground_residential", "Ground Residential"),
        ("2da", "2nd Day Air"),
        ("nda", "Next Day Air"),
    ]:
        svc_rsm = rsm.get(svc_key, {})
        if not svc_rsm:
            continue

        # Our sell base rates
        ws = wb.add_worksheet(f"{svc_label[:25]} Rates")
        ws.set_column("A:A", 12)
        ws.set_column("B:I", 12)

        markup_pct = summary.get("markup_pct", 0)

        ws.write(0, 0, "Weight", header_fmt)
        for j, z in enumerate(ZONES):
            ws.write(0, j + 1, f"Zone {z}", header_fmt)

        weights_in_rsm = sorted(int(w) for w in svc_rsm.keys())
        for i, w in enumerate(weights_in_rsm):
            ws.write(i + 1, 0, w, plain_fmt)
            for j, z in enumerate(ZONES):
                cell = svc_rsm.get(str(w), {}).get(str(z))
                if cell:
                    sell = cell["rate"] * (1 + markup_pct / 100)
                    ws.write(i + 1, j + 1, sell, green_fmt)
                else:
                    ws.write(i + 1, j + 1, "—", dash_fmt)

    # ── PACKAGE COMPARISON TAB ────────────────────────────────────────────────
    ws = wb.add_worksheet("Package Comparison")
    ws.set_column("A:A", 20)
    ws.set_column("B:M", 14)
    ws.freeze_panes(1, 0)

    headers = [
        "Tracking", "Weight", "Zone", "Service",
        "Their Base", "Their Fuel", "Their Accessorials", "Their Total",
        "Our Sell Base", "Our Sell Total",
        "Our Cost", "Savings", "Status",
    ]
    for j, h in enumerate(headers):
        ws.write(0, j, h, header_fmt)

    active = [p for p in packages if not p.get("correction_flag")]
    for i, pkg in enumerate(active):
        status = pkg.get("status", "no_rate")
        row_fmt = {
            "green": green_fmt, "grey": grey_fmt,
        }.get(status, money_fmt)

        their_acc = (
            (pkg.get("das_charge") or 0) +
            (pkg.get("das_extended") or 0) +
            (pkg.get("residential_charge") or 0) +
            (pkg.get("address_correction") or 0) +
            (pkg.get("additional_handling") or 0)
        )
        savings = pkg.get("customer_savings", 0) if status == "green" else 0.0

        ws.write(i + 1, 0, pkg.get("tracking", ""), plain_fmt)
        ws.write(i + 1, 1, pkg.get("weight"), plain_fmt)
        ws.write(i + 1, 2, pkg.get("zone", ""), plain_fmt)
        ws.write(i + 1, 3, pkg.get("service", ""), plain_fmt)
        ws.write(i + 1, 4, pkg.get("base_charge") or 0, money_fmt)
        ws.write(i + 1, 5, pkg.get("fuel_surcharge") or 0, money_fmt)
        ws.write(i + 1, 6, their_acc, money_fmt)
        ws.write(i + 1, 7, pkg.get("customer_total") or 0, money_fmt)
        sell_base = pkg.get("sell_base")
        sell_total = pkg.get("sell_total")
        ws.write(i + 1, 8, sell_base if sell_base is not None else "N/A", row_fmt if sell_base else plain_fmt)
        ws.write(i + 1, 9, sell_total if sell_total is not None else "N/A", row_fmt if sell_total else plain_fmt)
        ws.write(i + 1, 10, pkg.get("our_cost_total") or 0, money_fmt)
        ws.write(i + 1, 11, savings, green_fmt if savings > 0 else grey_fmt)
        ws.write(i + 1, 12, status.upper(), plain_fmt)

    # Totals row
    n = len(active) + 1
    ws.write(n, 0, "TOTALS", label_fmt)
    for col_idx, field in [
        (4, "base_charge"), (5, "fuel_surcharge"), (7, "customer_total"),
        (9, "sell_total"), (10, "our_cost_total"),
    ]:
        total = sum((p.get(field) or 0) for p in active)
        ws.write(n, col_idx, total, money_fmt)
    savings_total = sum(
        (p.get("customer_savings") or 0) for p in active if p.get("status") == "green"
    )
    ws.write(n, 11, savings_total, green_fmt)

    # ── ACCESSORIALS TAB ─────────────────────────────────────────────────────
    ws = wb.add_worksheet("Accessorials")
    ws.set_column("A:C", 25)
    ws.write(0, 0, "Accessorial", header_fmt)
    ws.write(0, 1, "Their Rate", header_fmt)
    ws.write(0, 2, "Our Rate", header_fmt)

    acc_rows = [
        ("Delivery Area Surcharge (DAS)", "Per carrier schedule", "Per carrier schedule"),
        ("DAS Extended", "Per carrier schedule", "Per carrier schedule"),
        ("Residential Delivery", "Per carrier schedule", "Per carrier schedule"),
        ("Address Correction", "Per carrier schedule", "Per carrier schedule"),
        ("Additional Handling", "Per carrier schedule", "Per carrier schedule"),
        ("Fuel Surcharge", f'{summary.get("customer_fuel_pct",0):.1f}%', f'{summary.get("our_fuel_pct",0):.1f}%'),
    ]
    for i, (name, their, ours) in enumerate(acc_rows, start=1):
        ws.write(i, 0, name, plain_fmt)
        ws.write(i, 1, their, plain_fmt)
        ws.write(i, 2, ours, plain_fmt)

    wb.close()
    return output.getvalue()


def generate_internal_excel(scenario_result: dict, scenario_name: str = "Internal") -> bytes:
    output = io.BytesIO()
    wb = xlsxwriter.Workbook(output, {"in_memory": True})

    summary = scenario_result["summary"]
    packages = scenario_result["packages"]
    rsm = scenario_result.get("rate_source_maps", {})

    header_fmt = wb.add_format({
        "bold": True, "bg_color": HEADER_BG, "font_color": HEADER_FONT,
        "border": 1, "align": "center",
    })
    money_fmt = wb.add_format({"num_format": "$#,##0.00", "border": 1, "align": "right"})
    green_fmt = wb.add_format({
        "bg_color": GREEN, "font_color": GREEN_FONT,
        "num_format": "$#,##0.00", "border": 1, "align": "right",
    })
    grey_fmt = wb.add_format({
        "bg_color": GREY, "font_color": GREY_FONT,
        "num_format": "$#,##0.00", "border": 1, "align": "right",
    })
    red_fmt = wb.add_format({
        "bg_color": RED, "font_color": RED_FONT,
        "num_format": "$#,##0.00", "border": 1, "align": "right",
    })
    plain_fmt = wb.add_format({"border": 1, "align": "center"})
    label_fmt = wb.add_format({"bold": True})
    title_fmt = wb.add_format({"bold": True, "font_size": 14})

    # ── EXECUTIVE SUMMARY ────────────────────────────────────────────────────
    ws = wb.add_worksheet("Executive Summary")
    ws.set_column("A:A", 35)
    ws.set_column("B:B", 20)
    ws.write("A1", f"Internal — {scenario_name}", title_fmt)

    rows = [
        ("Total Packages (active)", summary.get("total_packages")),
        ("Green (we win + save customer)", summary.get("green_count")),
        ("Grey (match, keep spread)", summary.get("grey_count")),
        ("Red (can't compete)", summary.get("red_count")),
        ("No Rate Available", summary.get("no_rate_count")),
        ("Corrections Excluded", summary.get("corrections_count")),
        ("", ""),
        ("Customer Current Monthly", f'${summary.get("customer_total_monthly",0):,.2f}'),
        ("Our Sell Monthly", f'${summary.get("sell_total_monthly",0):,.2f}'),
        ("Our Cost Monthly", f'${summary.get("our_cost_total_monthly",0):,.2f}'),
        ("Customer Savings Monthly", f'${summary.get("customer_savings_monthly",0):,.2f}'),
        ("Odyssey Margin Monthly", f'${summary.get("margin_monthly",0):,.2f}'),
        ("Margin %", f'{summary.get("margin_pct",0):.2f}%'),
        ("", ""),
        ("Markup %", f'{summary.get("markup_pct",0):.1f}%'),
        ("Our Fuel %", f'{summary.get("our_fuel_pct",0):.1f}%'),
        ("Customer Fuel %", f'{summary.get("customer_fuel_pct",0):.1f}%'),
        ("Fuel Spread", f'{summary.get("fuel_spread_pct",0):.1f}%'),
        ("Volume Multiplier", summary.get("volume_multiplier", 1.0)),
        ("", ""),
        ("BILLING INSTRUCTIONS", ""),
        ("Bill customer sell base + customer fuel % + pass-through accessorials", ""),
        ("Rate source: see Rate Source Map tab", ""),
        ("Fuel: apply customer fuel % to sell base each invoice", ""),
    ]
    for i, (label, val) in enumerate(rows, start=2):
        ws.write(f"A{i}", label, label_fmt if label else wb.add_format())
        ws.write(f"B{i}", val)

    # ── RATE SOURCE MAP ───────────────────────────────────────────────────────
    for svc_key, svc_label in [
        ("ground_commercial", "GC Source Map"),
        ("ground_residential", "GR Source Map"),
        ("2da", "2DA Source Map"),
        ("nda", "NDA Source Map"),
    ]:
        svc_rsm = rsm.get(svc_key, {})
        if not svc_rsm:
            continue
        ws = wb.add_worksheet(svc_label)
        ws.set_column("A:A", 10)
        ws.set_column("B:I", 22)
        ws.write(0, 0, "Weight", header_fmt)
        for j, z in enumerate(ZONES):
            ws.write(0, j + 1, f"Zone {z}", header_fmt)

        weights_in_rsm = sorted(int(w) for w in svc_rsm.keys())
        for i, w in enumerate(weights_in_rsm):
            ws.write(i + 1, 0, w, plain_fmt)
            for j, z in enumerate(ZONES):
                cell = svc_rsm.get(str(w), {}).get(str(z))
                if cell:
                    ws.write(i + 1, j + 1, cell.get("carrier_name", ""), plain_fmt)
                else:
                    ws.write(i + 1, j + 1, "—", plain_fmt)

    # ── COST & SELL GRIDS per service ─────────────────────────────────────────
    for svc_key, svc_label in [
        ("ground_commercial", "GC"),
        ("ground_residential", "GR"),
        ("2da", "2DA"),
        ("nda", "NDA"),
    ]:
        svc_rsm = rsm.get(svc_key, {})
        if not svc_rsm:
            continue
        markup_pct = summary.get("markup_pct", 0)
        our_fuel = summary.get("our_fuel_pct", 0)
        cust_fuel = summary.get("customer_fuel_pct", 0)

        for tab_label, rate_fn, fmt in [
            (f"{svc_label} Best Cost",
             lambda r, mup=markup_pct, of=our_fuel: r * (1 + of / 100),
             money_fmt),
            (f"{svc_label} Sell Base",
             lambda r, mup=markup_pct: r * (1 + mup / 100),
             green_fmt),
        ]:
            ws = wb.add_worksheet(tab_label[:31])
            ws.set_column("A:A", 10)
            ws.set_column("B:I", 14)
            ws.write(0, 0, "Weight", header_fmt)
            for j, z in enumerate(ZONES):
                ws.write(0, j + 1, f"Zone {z}", header_fmt)
            weights_in_rsm = sorted(int(w) for w in svc_rsm.keys())
            for i, w in enumerate(weights_in_rsm):
                ws.write(i + 1, 0, w, plain_fmt)
                for j, z in enumerate(ZONES):
                    cell = svc_rsm.get(str(w), {}).get(str(z))
                    if cell:
                        ws.write(i + 1, j + 1, rate_fn(cell["rate"]), fmt)
                    else:
                        ws.write(i + 1, j + 1, "—", plain_fmt)

    # ── MARGIN PER PACKAGE GRID ────────────────────────────────────────────────
    ws = wb.add_worksheet("Pkg P&L Detail")
    ws.set_column("A:A", 20)
    ws.set_column("B:P", 13)
    ws.freeze_panes(1, 0)

    hdrs = [
        "Tracking", "Weight", "Zone", "Service", "Correction?",
        "Their Base", "Their Fuel", "Their Total",
        "Best Rate", "Carrier", "Our Cost Base", "Our Cost Total",
        "Sell Base", "Sell Fuel", "Sell Total", "Margin", "Status",
    ]
    for j, h in enumerate(hdrs):
        ws.write(0, j, h, header_fmt)

    for i, pkg in enumerate(packages):
        status = pkg.get("status", "no_rate")
        row_color = {"green": green_fmt, "grey": grey_fmt, "red": red_fmt}.get(status, plain_fmt)

        ws.write(i + 1, 0, pkg.get("tracking", ""), plain_fmt)
        ws.write(i + 1, 1, pkg.get("weight"), plain_fmt)
        ws.write(i + 1, 2, pkg.get("zone", ""), plain_fmt)
        ws.write(i + 1, 3, pkg.get("service", ""), plain_fmt)
        ws.write(i + 1, 4, "YES" if pkg.get("correction_flag") else "no", plain_fmt)
        ws.write(i + 1, 5, pkg.get("base_charge") or 0, money_fmt)
        ws.write(i + 1, 6, pkg.get("fuel_surcharge") or 0, money_fmt)
        ws.write(i + 1, 7, pkg.get("customer_total") or 0, money_fmt)
        br = pkg.get("best_rate")
        ws.write(i + 1, 8, br if br is not None else "N/A", money_fmt if br else plain_fmt)
        ws.write(i + 1, 9, pkg.get("carrier_name", "") or "", plain_fmt)
        oct = pkg.get("our_cost_total")
        ws.write(i + 1, 10, pkg.get("our_cost_base") or 0, money_fmt)
        ws.write(i + 1, 11, oct if oct is not None else "N/A", money_fmt if oct else plain_fmt)
        sb = pkg.get("sell_base")
        st = pkg.get("sell_total")
        ws.write(i + 1, 12, sb if sb is not None else "N/A", row_color if sb else plain_fmt)
        ws.write(i + 1, 13, pkg.get("sell_fuel") or 0, money_fmt)
        ws.write(i + 1, 14, st if st is not None else "N/A", row_color if st else plain_fmt)
        m = pkg.get("margin")
        ws.write(i + 1, 15, m if m is not None else "N/A", row_color if m else plain_fmt)
        ws.write(i + 1, 16, status.upper(), plain_fmt)

    wb.close()
    return output.getvalue()


def generate_deal_brief(scenario_result: dict, account_name: str = "Account",
                         strategy_notes: str = "",
                         risks: str = "",
                         scenarios_comparison: Optional[list] = None) -> bytes:
    summary = scenario_result["summary"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def heading(text, level=1):
        h = doc.add_heading(text, level=level)
        h.style.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)

    def add_table_row(tbl, cells, bold_first=False):
        row = tbl.add_row()
        for i, (cell_obj, val) in enumerate(zip(row.cells, cells)):
            cell_obj.text = str(val)
            if bold_first and i == 0:
                cell_obj.paragraphs[0].runs[0].bold = True

    doc.add_heading(f"Deal Brief — {account_name}", 0)
    doc.add_paragraph(f"Prepared by: Odyssey Logistics | Date: " +
                      "{{today}}")

    heading("1. Account Overview")
    doc.add_paragraph(
        f"Account: {account_name}\n"
        f"Monthly Packages: {summary.get('total_packages', 0):,}\n"
        f"Current Monthly Spend: ${summary.get('customer_total_monthly', 0):,.2f}"
    )

    if strategy_notes:
        heading("2. Strategy")
        doc.add_paragraph(strategy_notes)
    else:
        heading("2. Strategy")
        doc.add_paragraph(
            "Primary objective: win the account with meaningful savings while maintaining "
            "target margin. Green packages represent immediate wins where the customer saves "
            "money. Grey packages are matched at current price — we earn the fuel spread. "
            "Red packages are passed or flagged for further negotiation."
        )

    heading("3. Monthly Economics")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    tbl.rows[0].cells[0].text = "Metric"
    tbl.rows[0].cells[1].text = "Value"
    for label, val in [
        ("Customer Current Monthly", f'${summary.get("customer_total_monthly",0):,.2f}'),
        ("Proposed Sell Monthly", f'${summary.get("sell_total_monthly",0):,.2f}'),
        ("Customer Savings Monthly", f'${summary.get("customer_savings_monthly",0):,.2f}'),
        ("Odyssey Cost Monthly", f'${summary.get("our_cost_total_monthly",0):,.2f}'),
        ("Odyssey Margin Monthly", f'${summary.get("margin_monthly",0):,.2f}'),
        ("Margin %", f'{summary.get("margin_pct",0):.2f}%'),
        ("Green Packages", str(summary.get("green_count",0))),
        ("Grey Packages", str(summary.get("grey_count",0))),
        ("Red Packages", str(summary.get("red_count",0))),
    ]:
        add_table_row(tbl, [label, val], bold_first=True)

    heading("4. Margin Sources")
    doc.add_paragraph(
        f"Base markup: {summary.get('markup_pct',0):.1f}% applied to best available rate.\n"
        f"Fuel spread: Customer pays {summary.get('customer_fuel_pct',0):.1f}% fuel; "
        f"we pay {summary.get('our_fuel_pct',0):.1f}% fuel — "
        f"{summary.get('fuel_spread_pct',0):.1f}% spread on every dollar of base.\n"
        f"Accessorials: passed through at invoice rate (see Accessorials tab in proposal)."
    )

    heading("5. Pricing Mechanics")
    doc.add_paragraph(
        "Sell base = best carrier rate × (1 + markup%)\n"
        "Sell fuel = sell base × customer fuel%\n"
        "Our cost = best carrier rate × (1 + our fuel%)\n"
        "Margin per package = sell total − our cost total\n"
        "Status: GREEN if sell total < their current total; "
        "GREY if our cost < their total (match their price); "
        "RED if our cost ≥ their total."
    )

    if scenarios_comparison:
        heading("6. Scenario Comparison")
        cols = ["Scenario", "Markup", "Green", "Grey", "Red",
                "Margin/Mo", "Cust Savings/Mo"]
        tbl = doc.add_table(rows=1, cols=len(cols))
        tbl.style = "Table Grid"
        for i, h in enumerate(cols):
            tbl.rows[0].cells[i].text = h
        for sc in scenarios_comparison:
            s = sc.get("summary", {})
            add_table_row(tbl, [
                sc.get("name", ""),
                f'{s.get("markup_pct",0):.1f}%',
                s.get("green_count", 0),
                s.get("grey_count", 0),
                s.get("red_count", 0),
                f'${s.get("margin_monthly",0):,.2f}',
                f'${s.get("customer_savings_monthly",0):,.2f}',
            ])

    heading("7. Billing Workflow")
    doc.add_paragraph(
        "1. Receive carrier invoice weekly.\n"
        "2. Match each package to rate source map (see internal workbook).\n"
        "3. Bill customer: sell base + (sell base × customer fuel%) + pass-through accessorials.\n"
        "4. Verify margin on each invoice before sending.\n"
        "5. Flag any correction or rebill rows — exclude from savings calculations."
    )

    heading("8. Risks")
    if risks:
        doc.add_paragraph(risks)
    else:
        doc.add_paragraph(
            "• Zone chart changes by carrier — re-validate quarterly.\n"
            "• Fuel surcharge fluctuates — lock customer fuel % or tie to published index.\n"
            "• DAS / residential mix can shift — monitor monthly.\n"
            "• Red packages: do not promise savings; present as pricing parity."
        )

    heading("9. Pitch Script")
    doc.add_paragraph(
        f"\"Based on your actual shipping data, we can save you approximately "
        f"${summary.get('customer_savings_monthly',0):,.0f}/month — "
        f"${summary.get('customer_savings_monthly',0)*12:,.0f} annualized — "
        f"with zero disruption to your current operations. "
        f"You keep your carrier relationships; we optimize the economics. "
        f"Every number in this proposal is auditable back to your own invoice data.\""
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
